"""PDF and Word renderer for the versioned print document contract."""
import hashlib
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from fontTools.ttLib import TTFont as GlyphFont
from PIL import Image as PillowImage
from reportlab.graphics import renderPDF
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Flowable,
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.doctemplate import LayoutError

from .contracts import (
    ExportError,
    FontSet,
    resolve_formula_image,
    snapshot_id,
    validate_document,
)
from .math import formula_strings, metric, omath, paint
from .rich import RichBreak, RichRun, parse_rich, pdf_markup, visible_runs


_PALETTE = {
    "key": ("175A8E", "EAF3FA"),
    "warn": ("8C4B00", "FFF3D6"),
    "bridge": ("60408A", "F1ECF8"),
    "erratum": ("A32535", "FCECEF"),
}
_FOOTER_FORMAT_GLYPHS = "0123456789 /|"
_FOOTER_STATUS_UNREVIEWED = "历史未审核"
_FORMULA_IMAGE_LABEL = "公式图片："
_FORMULA_SOURCE_LABEL = "公式来源："


def _font_digest(path):
    try:
        resolved = Path(path).resolve(strict=True)
    except OSError as exc:
        raise ExportError("font_not_found", "Every injected font path must exist") from exc
    if not resolved.is_file():
        raise ExportError("font_not_found", "Every injected font path must be a file")
    data = resolved.read_bytes()
    return resolved, hashlib.sha256(data).hexdigest()


def _register_fonts(fonts):
    roles = ("regular", "bold", "math")
    registered = {}
    cmap_by_role = {}
    for role in roles:
        path, digest = _font_digest(getattr(fonts, role))
        alias = f"StudyWorkbench_{role}_{digest}"
        registered[role] = alias
        if alias not in pdfmetrics.getRegisteredFontNames():
            try:
                pdfmetrics.registerFont(TTFont(alias, str(path)))
            except Exception as exc:
                raise ExportError("invalid_font", f"Could not register the injected {role} font") from exc
        try:
            face = GlyphFont(str(path), lazy=True)
            try:
                cmap_by_role[role] = face.getBestCmap() or {}
            finally:
                face.close()
        except Exception as exc:
            raise ExportError("invalid_font", f"Could not inspect the injected {role} font") from exc

    regular, bold = registered["regular"], registered["bold"]
    for family in {regular, bold}:
        pdfmetrics.registerFontFamily(
            family,
            normal=regular if family == regular else bold,
            bold=bold,
            italic=regular if family == regular else bold,
            boldItalic=bold,
        )
    return registered, cmap_by_role


def _add_glyphs(target, value):
    target.update(char for char in value if not char.isspace())


def _add_rich_glyphs(required, value, *, base_bold=False):
    for run in visible_runs(value):
        _add_glyphs(required["bold" if (base_bold or run.bold) else "regular"], run.text)


def _required_glyphs(document, snapshot):
    required = {"regular": set(), "bold": set(), "math": set()}
    _add_glyphs(required["regular"], document.title)
    _add_glyphs(required["regular"], _footer_status(document))
    _add_glyphs(required["regular"], snapshot[:12])
    _add_glyphs(required["regular"], _FOOTER_FORMAT_GLYPHS)
    for page in document.pages:
        for block in page:
            kind, content = block.kind, block.content
            if kind in {"title", "sub", "h", "p", "small", *_PALETTE}:
                _add_rich_glyphs(required, content, base_bold=kind in {"title", "h"})
            elif kind == "table":
                rows, _ = content
                for row in rows:
                    for value in row:
                        _add_rich_glyphs(required, value)
            elif kind == "math":
                for value in formula_strings(content):
                    _add_glyphs(required["math"], value)
            elif kind == "map":
                for line in content["root"]:
                    _add_glyphs(required["bold"], line)
                for group in content["groups"]:
                    for line in group["label"]:
                        _add_glyphs(required["bold"], line)
                    for line in group["detail"]:
                        _add_glyphs(required["regular"], line)
            elif kind == "formula_image":
                _add_glyphs(required["regular"], _FORMULA_IMAGE_LABEL + content["alt"])
                _add_glyphs(required["regular"], _FORMULA_SOURCE_LABEL + content["source_ref"])
    return required


def _check_glyphs(required, cmap_by_role):
    for role, characters in required.items():
        missing = sorted((char for char in characters if ord(char) not in cmap_by_role[role]), key=ord)
        if missing:
            point = f"U+{ord(missing[0]):04X}"
            raise ExportError("missing_glyph", f"The injected {role} font lacks required character {point}")


def _footer_status(document):
    if document.source.state == "legacy_unreviewed":
        return _FOOTER_STATUS_UNREVIEWED
    return document.source.state


def _paragraph_styles(fonts):
    regular, bold = fonts["regular"], fonts["bold"]
    styles = {
        "title": ParagraphStyle("a2_title", fontName=bold, fontSize=18, leading=25,
            spaceAfter=10, wordWrap="CJK", textColor=colors.HexColor("#163c65")),
        "sub": ParagraphStyle("a2_sub", fontName=regular, fontSize=9.5, leading=14,
            spaceAfter=9, wordWrap="CJK", textColor=colors.HexColor("#555555")),
        "h": ParagraphStyle("a2_h", fontName=bold, fontSize=12, leading=18,
            spaceBefore=8, spaceAfter=5, wordWrap="CJK", keepWithNext=True,
            textColor=colors.HexColor("#175a8e")),
        "p": ParagraphStyle("a2_p", fontName=regular, fontSize=11, leading=17.8,
            spaceAfter=7, wordWrap="CJK"),
        "small": ParagraphStyle("a2_small", fontName=regular, fontSize=9, leading=14,
            spaceAfter=5, wordWrap="CJK"),
        "cell": ParagraphStyle("a2_cell", fontName=regular, fontSize=9.4, leading=14.8,
            wordWrap="CJK"),
    }
    for kind, (foreground, background) in _PALETTE.items():
        styles[kind] = ParagraphStyle(
            f"a2_{kind}", parent=styles["p"], fontSize=10.5, leading=17,
            textColor=colors.HexColor(f"#{foreground}"), backColor=colors.HexColor(f"#{background}"),
            borderPadding=4, spaceBefore=3, spaceAfter=10,
        )
    return styles


def _box(drawing, x, y, width, height, lines, font_name, size, filled=False):
    drawing.add(Rect(
        x, y, width, height, rx=4, ry=4,
        fillColor=colors.HexColor("#eaf3fa" if filled else "#fafcfe"),
        strokeColor=colors.HexColor("#84a7c8"), strokeWidth=0.8,
    ))
    for index, line in enumerate(lines):
        drawing.add(String(
            x + width / 2,
            y + height / 2 + (len(lines) - 1) * 7 - index * 14 - 3.5,
            line,
            fontName=font_name,
            fontSize=size,
            textAnchor="middle",
            fillColor=colors.HexColor("#163c65" if filled else "#303b46"),
        ))


def _check_map_lines(lines, font_name, size, box_width):
    max_width = box_width - 12
    if any(pdfmetrics.stringWidth(line, font_name, size) > max_width for line in lines):
        raise ExportError("map_label_too_wide", "A map label exceeds its box width")


def _knowledge_map(content, fonts):
    root = content["root"]
    groups = content["groups"]
    _check_map_lines(root, fonts["bold"], 10.7, 91)
    for group in groups:
        _check_map_lines(group["label"], fonts["bold"], 10.1, 116)
        _check_map_lines(group["detail"], fonts["regular"], 10, 253)
    drawing = Drawing(494, 378)
    for index, group in enumerate(groups):
        y = 319 - index * 62 - (6 - len(groups)) * 31
        drawing.add(Line(91, 188, 107, y + 26, strokeColor=colors.HexColor("#999999"), strokeWidth=0.7))
        drawing.add(Line(223, y + 26, 240, y + 26, strokeColor=colors.HexColor("#999999"), strokeWidth=0.7))
        _box(drawing, 107, y, 116, 52, group["label"], fonts["bold"], 10.1, True)
        _box(drawing, 240, y, 253, 52, group["detail"], fonts["regular"], 10)
    _box(drawing, 0, 157, 91, 62, root, fonts["bold"], 10.7, True)
    return drawing


class _MathLine(Flowable):
    def __init__(self, node, font_name):
        super().__init__()
        self.node = node
        self.font_name = font_name
        self.font_size = 11.6

    def wrap(self, width, height):
        natural = metric(self.node, self.font_name, 11.6)[0]
        inner_width = width - 12
        self.font_size = min(11.6, 11.6 * inner_width / natural) if natural else 11.6
        if self.font_size < 9.8:
            raise ExportError("formula_too_wide", "Formula cannot fit at the minimum supported print size")
        measured_width, above, below = metric(self.node, self.font_name, self.font_size)
        if measured_width + 12 > width + 0.1:
            raise ExportError("formula_too_wide", "Formula exceeds its printable width")
        self.width = measured_width + 12
        self.height = above + below + 11
        self.baseline = below + 5
        return self.width, self.height

    def draw(self):
        paint(self.canv, self.node, 7, self.baseline, self.font_size, self.font_name)


def _formula_image_flowables(content, path, styles):
    with PillowImage.open(path) as image:
        original_width, original_height = image.size
    width = float(content["width_points"])
    height = width * original_height / original_width
    if height > 700:
        raise ExportError("fallback_too_tall", "Formula image exceeds the supported printable height")
    return [
        Image(str(path), width=width, height=height),
        Paragraph(pdf_markup(_FORMULA_IMAGE_LABEL + content["alt"]), styles["small"]),
        Paragraph(pdf_markup(_FORMULA_SOURCE_LABEL + content["source_ref"]), styles["small"]),
    ]


def _pdf_table(content, styles):
    rows, widths = content
    table = Table(
        [[Paragraph(pdf_markup(value), styles["cell"]) for value in row] for row in rows],
        colWidths=widths,
        repeatRows=1,
        hAlign="LEFT",
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf3fa")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.HexColor("#84a7c8")),
        ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor("#d8e0e8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def _build_pdf(document, output_path, fonts, styles, snapshot, assets):
    story = []
    usable_width = A4[0] - 94
    for page_index, blocks in enumerate(document.pages):
        if page_index:
            story.append(PageBreak())
        for block in blocks:
            kind, content = block.kind, block.content
            if kind == "math":
                story.append(_MathLine(content, fonts["math"]))
            elif kind == "map":
                story.append(_knowledge_map(content, fonts))
            elif kind == "space":
                story.append(Spacer(1, float(content)))
            elif kind == "table":
                if sum(content[1]) > usable_width + 0.1:
                    raise ExportError("table_too_wide", "Table exceeds the A4 printable width")
                story.extend((_pdf_table(content, styles), Spacer(1, 6)))
            elif kind == "formula_image":
                path = assets[id(block)]
                story.extend(_formula_image_flowables(content, path, styles))
            else:
                story.append(Paragraph(pdf_markup(content), styles[kind]))

    actual_pages = []
    expected_pages = len(document.pages)
    status = _footer_status(document)
    footer_right = f"{status} | {snapshot[:12]} | {expected_pages}/{expected_pages}"
    if pdfmetrics.stringWidth(document.title, fonts["regular"], 8) + pdfmetrics.stringWidth(footer_right, fonts["regular"], 8) + 10 > A4[0] - 94:
        raise ExportError("footer_too_wide", "Document title and source reference do not fit the footer")

    def footer(canvas, template):
        actual_pages.append(template.page)
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#84a7c8"))
        canvas.setLineWidth(0.7)
        canvas.line(47, 35, A4[0] - 47, 35)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.setFont(fonts["regular"], 8)
        canvas.drawString(47, 23, document.title)
        canvas.drawRightString(A4[0] - 47, 23, f"{status} | {snapshot[:12]} | {template.page}/{expected_pages}")
        canvas.restoreState()

    template = SimpleDocTemplate(
        str(output_path), pagesize=A4, leftMargin=47, rightMargin=47,
        topMargin=42, bottomMargin=47, title=document.title, author="",
    )
    try:
        template.build(story, onFirstPage=footer, onLaterPages=footer)
    except LayoutError as exc:
        raise ExportError("layout_overflow", "PDF content exceeds the supported page layout") from exc
    if actual_pages != list(range(1, expected_pages + 1)):
        raise ExportError("page_count_mismatch", f"Expected {expected_pages} PDF pages; rendered {len(actual_pages)}")
    return len(actual_pages)


def _set_run_font(run, family):
    run.font.name = family
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), family)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), family)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), family)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), family)


def _word_rich(paragraph, value, family, *, color_override=None):
    for event in parse_rich(value):
        if isinstance(event, RichBreak):
            paragraph.add_run().add_break()
            continue
        run = paragraph.add_run(event.text)
        _set_run_font(run, family)
        run.bold = event.bold or run.bold
        run.font.superscript = event.superscript
        if color_override:
            run.font.color.rgb = RGBColor.from_string(color_override)
        elif event.color:
            run.font.color.rgb = RGBColor.from_string(event.color[1:])
        if event.backcolor:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), event.backcolor[1:])
            run._element.get_or_add_rPr().append(shading)


def _shade_paragraph(paragraph, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def _write_space(document, points):
    if points == 0:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(float(points))
    run = paragraph.add_run("\u00a0")
    _set_run_font(run, document.styles["Normal"].font.name)
    run.font.size = Pt(1)


def _write_formula_image(document, content, path, family):
    paragraph = document.add_paragraph()
    inline = paragraph.add_run().add_picture(str(path), width=Pt(float(content["width_points"])))
    inline._inline.docPr.set("descr", content["alt"])
    inline._inline.docPr.set("title", content["source_ref"])
    _word_rich(document.add_paragraph(), _FORMULA_IMAGE_LABEL + content["alt"], family)
    _word_rich(document.add_paragraph(), _FORMULA_SOURCE_LABEL + content["source_ref"], family)


def _word_table(document, content, family):
    rows, widths = content
    table = document.add_table(rows=0, cols=len(widths))
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Pt(width)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        row_properties = table.rows[-1]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(cells, values, widths):
            cell.text = ""
            cell.width = Pt(width)
            _word_rich(cell.paragraphs[0], value, family)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EAF3FA")
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True


def _build_docx(document, output_path, word_family, map_fonts, snapshot, assets, work_dir):
    word = Document()
    section = word.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = Mm(16.6)
    section.top_margin = Mm(14.8)
    section.bottom_margin = Mm(16.6)
    section.footer_distance = Mm(8)

    for name in ("Normal", "Title", "Heading 1"):
        style = word.styles[name]
        style.font.name = word_family
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), word_family)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), word_family)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), word_family)
        style._element.get_or_add_rPr().rFonts.set(qn("w:cs"), word_family)
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = word.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)
    title_style = word.styles["Title"]
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("163C65")
    title_style.paragraph_format.space_after = Pt(8)
    heading = word.styles["Heading 1"]
    heading.font.size = Pt(12)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    heading.font.color.rgb = RGBColor.from_string("175A8E")
    heading.font.bold = True

    map_index = 0
    for page_index, blocks in enumerate(document.pages):
        if page_index:
            word.add_page_break()
        for block in blocks:
            kind, content = block.kind, block.content
            if kind == "math":
                paragraph = word.add_paragraph()
                paragraph.paragraph_format.left_indent = Mm(3)
                paragraph.paragraph_format.keep_together = True
                equation = OxmlElement("m:oMath")
                for element in omath(content):
                    equation.append(element)
                paragraph._p.append(equation)
            elif kind == "map":
                temp_path = work_dir / f"map-{map_index}"
                map_index += 1
                try:
                    temp_path.mkdir()
                    map_pdf = temp_path / "map.pdf"
                    prefix = temp_path / "map"
                    renderPDF.drawToFile(_knowledge_map(content, map_fonts), str(map_pdf))
                    executable = shutil.which("pdftoppm")
                    if not executable:
                        raise ExportError("missing_renderer_tool", "pdftoppm is required to create the Word map image")
                    try:
                        result = subprocess.run(
                            [executable, "-f", "1", "-l", "1", "-singlefile", "-r", "216", "-png", str(map_pdf), str(prefix)],
                            capture_output=True, text=True, timeout=30, check=False,
                        )
                    except subprocess.TimeoutExpired as exc:
                        raise ExportError("map_render_timeout", "Map image conversion exceeded its time limit") from exc
                    if result.returncode != 0 or not prefix.with_suffix(".png").is_file():
                        raise ExportError("map_render_failed", "pdftoppm could not create the Word map image")
                    word.add_picture(str(prefix.with_suffix(".png")), width=Mm(174))
                except OSError as exc:
                    raise ExportError("map_render_failed", "Could not create the Word map image") from exc
            elif kind == "space":
                _write_space(word, content)
            elif kind == "table":
                _word_table(word, content, word_family)
            elif kind == "formula_image":
                path = assets[id(block)]
                _write_formula_image(word, content, path, word_family)
            else:
                style_name = "Title" if kind == "title" else "Heading 1" if kind == "h" else "Normal"
                paragraph = word.add_paragraph(style=style_name)
                color_override = None
                if kind in _PALETTE:
                    foreground, background = _PALETTE[kind]
                    color_override = foreground
                    _shade_paragraph(paragraph, background)
                _word_rich(paragraph, content, word_family, color_override=color_override)
                if kind in {"sub", "small"}:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                paragraph.paragraph_format.keep_together = True

    update_fields = word.settings._element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        word.settings._element.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    footer = section.footer.paragraphs[0]
    footer.add_run(document.title + " | " + _footer_status(document) + " | " + snapshot[:12] + " | ")
    for instruction in ("PAGE", "NUMPAGES"):
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        footer._p.append(field)
        if instruction == "PAGE":
            footer.add_run(" / ")
    for run in footer.runs:
        _set_run_font(run, word_family)
        run.font.size = Pt(8)
    word.core_properties.title = document.title
    word.core_properties.author = ""
    word.save(output_path)
    return sum(1 for page in document.pages for block in page if block.kind == "math")


def _stage_formula_images(document, asset_root, work_dir):
    assets = {}
    index = 0
    for page in document.pages:
        for block in page:
            if block.kind != "formula_image":
                continue
            source = resolve_formula_image(block.content, asset_root)
            raw = source.read_bytes()
            if hashlib.sha256(raw).hexdigest() != block.content["sha256"]:
                raise ExportError("fallback_hash_mismatch", "Formula image changed while it was being staged")
            staged = work_dir / f"formula-{index}.png"
            staged.write_bytes(raw)
            with PillowImage.open(staged) as image:
                image.load()
            assets[id(block)] = staged
            index += 1
    return assets


def render_document(document, output_dir: Path, fonts: FontSet, *, asset_root: Path | None = None) -> dict:
    """Render a validated document into one PDF and one editable Word file."""
    validate_document(document)
    snapshot = snapshot_id(document)
    registered, cmap_by_role = _register_fonts(fonts)
    _check_glyphs(_required_glyphs(document, snapshot), cmap_by_role)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = output_dir / "document.pdf"
    docx_path = output_dir / "document.docx"
    styles = _paragraph_styles(registered)
    with tempfile.TemporaryDirectory(prefix="study-workbench-render-", dir=output_dir) as temp_dir:
        work_dir = Path(temp_dir)
        assets = _stage_formula_images(document, asset_root, work_dir)
        page_count = _build_pdf(document, pdf_path, registered, styles, snapshot, assets)
        word_equations = _build_docx(document, docx_path, fonts.word_family, registered, snapshot, assets, work_dir)
    return {"pdf": pdf_path, "docx": docx_path, "page_count": page_count, "word_equations": word_equations}
