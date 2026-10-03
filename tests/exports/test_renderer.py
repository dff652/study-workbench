from __future__ import annotations

import hashlib
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from lxml import etree
from PIL import Image as PillowImage, ImageDraw

from app.exports import render_document
from app.exports.contracts import Block, ExportDocument, ExportError, SourceRef
from app.exports.fonts import prepare_fonts
from app.domain.arithmetic import formula_ast


MATH_FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf")
ROOT = Path(__file__).resolve().parents[2]


def source_ref():
    return SourceRef("synthetic-export-source", "a" * 64)


def document(pages, *, document_id="synthetic-render", title="Synthetic Export", purpose="knowledge_summary"):
    return ExportDocument(document_id, title, purpose, tuple(tuple(page) for page in pages), source_ref())


def pdf_text(path):
    result = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, timeout=15, check=False
    )
    if result.returncode:
        raise AssertionError(result.stderr)
    return result.stdout


def pdf_info(path):
    result = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, timeout=15, check=False)
    if result.returncode:
        raise AssertionError(result.stderr)
    return {key.strip(): value.strip() for key, value in (line.split(":", 1) for line in result.stdout.splitlines() if ":" in line)}


def pdf_page_count(path):
    try:
        return int(pdf_info(path)["Pages"])
    except KeyError as exc:
        raise AssertionError("pdfinfo did not report a page count") from exc


def word_xml(path):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("word/document.xml")), archive.namelist()


def font_coverage_document():
    formula = ["r", ["t", "x"], ["u", ["t", "n"], ["t", "2"]], ["f", ["t", "1"], ["t", "2"]], ["d", ["t", "x"], ["t", "i"]]]
    return document([[
        Block("title", "Synthetic Export", "title"),
        Block("sub", "Styled sample", "body"),
        Block("p", "bold phrase plain colored mark x⁴ second line Second page paragraph.", "body"),
        Block("math", formula, "body"),
        Block("table", ((
            ("Topic", "Detail"),
            ("Fraction", "Table value"),
        ), (220, 220)), "body"),
        Block("map", {"root": ["Study map", "Synthetic root"], "groups": [
            {"label": ["Group one"], "detail": ["First detail", "Second detail"]},
            {"label": ["Group two"], "detail": ["Third detail"]},
        ]}, "body"),
        Block("formula_image", {
            "storage_key": "synthetic.png", "sha256": "a" * 64,
            "source_ref": "synthetic-source-17", "alt": "handwritten radical alternative", "width_points": 120,
        }, "body"),
    ]])


class RendererTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.font_dir = tempfile.TemporaryDirectory(prefix="study-workbench-export-fonts-")
        cls.font_set, _ = prepare_fonts(
            [font_coverage_document()], Path(cls.font_dir.name) / "fonts",
            regular_source="/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            bold_source="/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
            math_source=MATH_FONT,
            cjk_notice=ROOT / "licenses" / "Noto-OFL.txt",
            math_notice=ROOT / "licenses" / "DejaVu-fonts.txt",
        )

    @classmethod
    def tearDownClass(cls):
        cls.font_dir.cleanup()

    def render(self, doc, folder, *, asset_root=None):
        return render_document(doc, folder, self.font_set, asset_root=asset_root)

    def test_pdf_and_word_preserve_two_pages_markup_table_map_math_and_writing_space(self):
        formula = [
            "r",
            ["t", "x"],
            ["u", ["t", "n"], ["t", "2"]],
            ["t", " + "],
            ["f", ["t", "1"], ["t", "2"]],
            ["t", " + "],
            ["d", ["t", "x"], ["t", "i"]],
        ]
        doc = document([
            [
                Block("title", "Synthetic Export", "title"),
                Block("sub", "Styled sample", "body"),
                Block("p", "<b>bold phrase</b> plain <font color=\"#A00000\" backColor=\"#FFF3D6\">colored mark</font> x⁴<br/>second line", "body"),
                Block("math", formula, "body"),
                Block("table", ([
                    ("Topic", "Detail"),
                    ("Fraction", "Table value"),
                ], (220, 220)), "body"),
                Block("space", 36, "body"),
            ],
            [
                Block("map", {
                    "root": ["Study map", "Synthetic root"],
                    "groups": [
                        {"label": ["Group one"], "detail": ["First detail", "Second detail"]},
                        {"label": ["Group two"], "detail": ["Third detail"]},
                    ],
                }, "body"),
                Block("p", "Second page paragraph.", "body"),
            ],
        ])
        with tempfile.TemporaryDirectory() as temp_dir:
            result = self.render(doc, Path(temp_dir))
            self.assertEqual(result["page_count"], 2)
            self.assertEqual(result["word_equations"], 1)
            self.assertEqual(result["pdf"].name, "document.pdf")
            self.assertEqual(result["docx"].name, "document.docx")
            self.assertEqual(pdf_page_count(result["pdf"]), 2)
            self.assertEqual(pdf_info(result["pdf"]).get("Title"), "Synthetic Export")
            extracted = pdf_text(result["pdf"])
            for visible in ("Synthetic Export", "bold phrase", "colored mark", "Fraction", "Table value", "Group one", "First detail", "Second page paragraph.", "历史未审核"):
                self.assertIn(visible, extracted)

            word = Document(result["docx"])
            paragraph_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
            table_text = "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
            self.assertIn("bold phrase", paragraph_text)
            self.assertIn("colored mark", paragraph_text)
            self.assertIn("Fraction", table_text)
            self.assertIn("Table value", table_text)
            self.assertEqual(len(word.tables), 1)
            self.assertTrue(any(p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY for p in word.paragraphs))

            xml_root, names = word_xml(result["docx"])
            ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
            self.assertEqual(len(xml_root.xpath(".//m:oMath", namespaces=ns)), 1)
            self.assertEqual(len(xml_root.xpath(".//m:f", namespaces=ns)), 1)
            self.assertEqual(len(xml_root.xpath(".//m:sSup", namespaces=ns)), 1)
            self.assertEqual(len(xml_root.xpath(".//m:sSub", namespaces=ns)), 1)
            self.assertTrue(any(name.startswith("word/media/") for name in names))
            footer_text = "\n".join(p.text for p in word.sections[0].footer.paragraphs)
            self.assertIn("历史未审核", footer_text)
            self.assertIn(" / ", footer_text)
            self.assertNotIn(" / 2", footer_text)
            with zipfile.ZipFile(result["docx"]) as archive:
                w_ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                footer_root = etree.fromstring(archive.read("word/footer1.xml"))
                settings_root = etree.fromstring(archive.read("word/settings.xml"))
            self.assertEqual(footer_root.xpath(".//w:fldSimple/@w:instr", namespaces=w_ns), ["PAGE", "NUMPAGES"])
            self.assertEqual(settings_root.xpath(".//w:updateFields/@w:val", namespaces=w_ns), ["true"])

    def test_formula_image_requires_sha_verified_local_png_and_has_visible_fallback(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            assets = root / "assets"
            assets.mkdir()
            image_path = assets / "formula.png"
            image = PillowImage.new("RGB", (160, 40), "white")
            ImageDraw.Draw(image).line((5, 30, 150, 5), fill="black", width=3)
            image.save(image_path, format="PNG")
            content = {
                "storage_key": "formula.png",
                "sha256": hashlib.sha256(image_path.read_bytes()).hexdigest(),
                "source_ref": "synthetic-source-17",
                "alt": "handwritten radical alternative",
                "width_points": 120,
            }
            doc = document([[Block("formula_image", content, "body")]])
            output = root / "out"
            result = self.render(doc, output, asset_root=assets)
            extracted = pdf_text(result["pdf"])
            self.assertIn("公式图片：handwritten radical alternative", extracted)
            self.assertIn("公式来源：synthetic-source-17", extracted)
            xml_root, _ = word_xml(result["docx"])
            ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
            self.assertEqual(xml_root.xpath(".//wp:docPr/@descr", namespaces=ns), ["handwritten radical alternative"])
            self.assertEqual(xml_root.xpath(".//wp:docPr/@title", namespaces=ns), ["synthetic-source-17"])

            bad_content = dict(content, sha256="b" * 64)
            bad_doc = document([[Block("formula_image", bad_content, "body")]], document_id="wrong-hash")
            bad_output = root / "bad-output"
            with self.assertRaises(ExportError) as caught:
                self.render(bad_doc, bad_output, asset_root=assets)
            self.assertEqual(caught.exception.code, "fallback_hash_mismatch")
            self.assertEqual(list(bad_output.iterdir()), [])

    def test_power_base_grouping_is_visible_in_pdf_and_word_math(self):
        expressions = ("(-2)^2", "(-x)^2", "(x^2)^3", "-x^2")
        with tempfile.TemporaryDirectory() as temp_dir:
            for index, expression in enumerate(expressions):
                output = Path(temp_dir) / f"power-{index}"
                doc = document([[Block("math", formula_ast(expression), "answer")]],
                    document_id=f"power-{index}")
                result = self.render(doc, output)
                text = pdf_text(result["pdf"]).replace(" ", "")
                expected = {
                    "(-2)^2": "(-2)",
                    "(-x)^2": "(-x)",
                    "(x^2)^3": "(x)",
                    "-x^2": "-x",
                }[expression]
                self.assertIn(expected, text)
                xml_root, _ = word_xml(result["docx"])
                ns = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
                math_text = "".join(xml_root.xpath(".//m:t/text()", namespaces=ns))
                self.assertEqual(math_text, {
                    "(-2)^2": "(-2)2",
                    "(-x)^2": "(-x)2",
                    "(x^2)^3": "(x2)3",
                    "-x^2": "-x2",
                }[expression])

    def test_missing_glyph_fails_before_creating_output(self):
        doc = document([[Block("p", "unsupported \U0010ffff", "body")]], document_id="missing-glyph")
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "not-created"
            with self.assertRaises(ExportError) as caught:
                self.render(doc, output)
            self.assertEqual(caught.exception.code, "missing_glyph")
            self.assertFalse(output.exists())

    def test_unsupported_formula_and_markup_fail_closed(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            for name, block in (
                ("unknown-formula", Block("math", ["sqrt", ["t", "x"]], "body")),
                ("unknown-markup", Block("p", "<script>not allowed</script>", "body")),
                ("processing-instruction", Block("p", "before <?hide content?> after", "body")),
                ("unknown-declaration", Block("p", "<![CDATA[hidden content]]>", "body")),
            ):
                output = Path(temp_dir) / name
                with self.subTest(name=name), self.assertRaises(ExportError):
                    self.render(document([[block]], document_id=name), output)
                self.assertFalse(output.exists())

    def test_knowledge_map_rejects_labels_that_exceed_box_width(self):
        doc = document([[Block("map", {
            "root": ["X" * 100],
            "groups": [{"label": ["Group"], "detail": ["Detail"]}],
        }, "body")]], document_id="wide-map-label")
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "wide-map"
            with self.assertRaises(ExportError) as caught:
                self.render(doc, output)
            self.assertEqual(caught.exception.code, "map_label_too_wide")
            self.assertEqual(list(output.iterdir()), [])

    def test_import_has_no_font_registration_side_effect(self):
        code = """
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as ReportLabTTFont
from fontTools.ttLib import TTFont as FontToolsTTFont
before = set(pdfmetrics.getRegisteredFontNames())
def fail_on_font_open(*args, **kwargs):
    raise AssertionError('font opened during import')
ReportLabTTFont.__init__ = fail_on_font_open
FontToolsTTFont.__init__ = fail_on_font_open
import app.exports.renderer
after = set(pdfmetrics.getRegisteredFontNames())
assert before == after, (before, after)
"""
        result = subprocess.run(
            [sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True, timeout=15, check=False
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_sequential_exports_with_distinct_cjk_subsets_keep_all_text(self):
        # Each business export prepares a different subset in the same web
        # process. ReportLab must not reuse another subset with the same face.
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            for index, text in enumerate(("甲乙丙", "丁戊己")):
                doc = document([[Block("p", text, "body")]],
                    document_id=f"distinct-subset-{index}")
                fonts, _ = prepare_fonts([doc], root / f"fonts-{index}",
                    regular_source="/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                    bold_source="/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
                    math_source=MATH_FONT,
                    cjk_notice=ROOT / "licenses" / "Noto-OFL.txt",
                    math_notice=ROOT / "licenses" / "DejaVu-fonts.txt")
                result = render_document(doc, root / f"output-{index}", fonts)
                extracted = subprocess.run(["pdftotext", str(result["pdf"]), "-"],
                    check=True, capture_output=True, text=True).stdout
                self.assertIn(text, extracted)


if __name__ == "__main__":
    unittest.main()
